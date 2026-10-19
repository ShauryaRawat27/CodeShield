from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("/Users/shauryarawat/Documents/ChatGPT/fog/outputs/codeshield-deliverables/CodeShield_Review3_Report_Prompt_Pack.docx")

NAVY = "173B57"
BLUE = "2F6F8F"
PALE_BLUE = "EAF3F7"
PALE_GRAY = "F4F6F7"
MID_GRAY = "D9D9D9"
TEXT = "1F2933"
MUTED = "53606B"


def set_font(run, name="Arial", size=10.5, bold=None, italic=None, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.extend([r_fonts, color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_para(doc, text="", bold_lead=None, style=None, italic=False, align=None, space_after=6):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_font(first, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest, italic=italic)
    else:
        run = p.add_run(text)
        set_font(run, italic=italic)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        set_font(p.add_run(item))


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.27)
        p.paragraph_format.first_line_indent = Inches(-0.27)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.1
        set_font(p.add_run(f"{index}.  {item}"))


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_table_header(header)
    for index, title in enumerate(headers):
        cell = header.cells[index]
        shade_cell(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 110, 110, 110, 110)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(title), size=9.2, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if row_index % 2 == 1:
            for cell in row.cells:
                shade_cell(cell, PALE_GRAY)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(str(value)), size=font_size)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor.from_string(TEXT)
styles["Title"].font.name = "Arial"
styles["Title"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Title"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Title"].font.size = Pt(25)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
title_ppr = styles["Title"].element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

doc.core_properties.title = "CodeShield Review 3 Report Writing Brief"
doc.core_properties.subject = "Verified project information and master prompt for Review 3"
doc.core_properties.author = ""
doc.core_properties.keywords = "CodeShield, security drift, Review 3, research report"

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(footer_p.add_run("CodeShield Review 3 Report Writing Brief"), size=9, color=MUTED)
tab = footer_p.add_run("\t")
set_font(tab, size=9)
add_page_number(footer.add_paragraph())

# Cover
for _ in range(3):
    doc.add_paragraph()
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(14)
title.add_run("CodeShield Review 3 Report Writing Brief")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(22)
set_font(subtitle.add_run("Verified project facts section guidance experimental evidence and a reusable master prompt"), size=13, color=MUTED)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_after = Pt(8)
set_font(meta.add_run("Project  CodeShield Context Aware Security Drift Analyzer"), size=11, bold=True)
date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(date_p.add_run("Prepared for Review 3  October 2026"), size=10.5, color=MUTED)

doc.add_page_break()

add_heading(doc, "How to use this document", 1)
add_para(doc, "This document contains the factual material needed to write a Review 3 report with the sections Abstract, Introduction, Relevant Work, Methodology, Experimental Results, and Conclusions. It is a writing brief and prompt pack rather than the final report. Copy the master prompt at the end into a writing tool, attach the listed project files when possible, and retain the truth boundaries in this document.")
add_para(doc, "Main conclusion  CodeShield is a functional Spring Boot prototype for deterministic security drift detection. It compares a registered application baseline with repository configuration or supplied runtime observations, produces evidence linked findings, assigns context adjusted risk scores, and stores scan history. The current implementation supports real public Git repositories, but it does not yet include autonomous runtime agents, live vulnerability feeds, authentication, or machine learning.")

add_heading(doc, "Recommended report scope", 2)
add_table(doc, ["Section", "Recommended length", "Purpose"], [
    ["Abstract", "180 to 250 words", "Summarize the problem, method, implementation, verified results, and conclusion."],
    ["Introduction", "800 to 1100 words", "Define security drift, motivate the problem, state the gap, objectives, contributions, and scope."],
    ["Relevant Work", "900 to 1200 words", "Synthesize recent container, configuration, dependency, SBOM, exposure, and runtime research."],
    ["Methodology", "1200 to 1600 words", "Explain the design, data model, repository inspection, rule engine, scoring, persistence, and evaluation plan."],
    ["Experimental Results", "700 to 1000 words", "Report only measured tests, benchmark composition, and recorded repository scans with limitations."],
    ["Conclusions", "350 to 500 words", "Answer the objectives, summarize evidence, state limitations, and identify specific next work."],
], widths=[1.25, 1.35, 4.65])

add_heading(doc, "Verified project fact sheet", 1)
add_table(doc, ["Area", "Verified information"], [
    ["Project title", "CodeShield Context Aware Security Drift Analyzer"],
    ["Problem", "Security controls can drift after deployment. Runtime versions, dependency versions, network exposure, container users, mounts, and image identities may differ from the approved state."],
    ["Primary objective", "Detect and explain mismatches between application intent and observed repository or runtime evidence."],
    ["Implementation", "Java 21, Spring Boot 3.3.6, Spring MVC, Thymeleaf, Spring Data JPA, H2, Maven, JUnit 5, and Eclipse JGit."],
    ["Repository inputs", "Public HTTPS Git clone URLs or readable local repository paths supplied by the user."],
    ["Supported files", "pom.xml, build.gradle, build.gradle.kts, Dockerfile variants, compose.yaml, compose.yml, docker-compose.yaml, and docker-compose.yml."],
    ["Runtime inputs", "Manual browser form or POST /api/scans. Values include Java version, database exposure, container user, Docker socket mount, image digest, dependency version, Internet exposure, exploit evidence, and evidence source."],
    ["Output", "Persisted findings with a rule ID, expected state, observed state, evidence, remediation, score, severity, and workflow status."],
    ["Persistence", "Embedded H2 file database in the current implementation."],
    ["AI status", "No AI or machine learning is used. Detection and scoring are deterministic and explainable."],
], widths=[1.55, 5.7])

add_heading(doc, "Implemented scope and future scope", 2)
add_table(doc, ["Implemented now", "Future work and not current claims"], [
    ["Register application baselines and ownership", "User authentication, Spring Security, and role based access control"],
    ["Inspect public Git repositories and local directories", "Private repository authentication and organization wide repository discovery"],
    ["Parse Maven, Gradle, Dockerfile, and Compose configuration", "Kubernetes manifests, Terraform, SBOM ingestion, and additional language ecosystems"],
    ["Accept manual and REST API runtime observations", "Read only Docker Engine collector and authenticated host agent"],
    ["Apply six deterministic drift rules", "CVE, OSV, NVD, and CISA KEV enrichment"],
    ["Persist findings and scan history in H2", "PostgreSQL, durable queues, scheduled distributed scans, and alert integrations"],
    ["Track finding status as open, resolved, or accepted", "Automated remediation and policy orchestration"],
], widths=[3.6, 3.65])

doc.add_page_break()

add_heading(doc, "Abstract content guide", 1)
add_para(doc, "The abstract should be written as one compact paragraph. It should not contain citations unless the required institutional format explicitly allows them.")
add_numbered(doc, [
    "Open with the deployment drift problem: an approved source or configuration state may differ from what a repository or runtime observation shows later.",
    "State the objective: develop an explainable system that compares application intent with observed evidence and prioritizes mismatches.",
    "Describe the method: registered baselines, repository configuration extraction, normalized scan input, six deterministic rules, and context based scoring.",
    "Name the implementation: Java 21 and Spring Boot with a web interface, REST API, persistence, JGit repository acquisition, and automated tests.",
    "Report verified evidence: 11 of 11 tests passed; the benchmark contains 120 labeled cases; two public repository scans completed and produced five and two findings respectively.",
    "Close carefully: the prototype demonstrates feasible evidence linked drift detection, while broader runtime collection and independently measured precision and recall remain future work.",
])
add_para(doc, "Do not state that the system uses artificial intelligence, continuously monitors live infrastructure, automatically discovers vulnerabilities, or has achieved a measured accuracy percentage. None of those claims has been established.", italic=True)

add_heading(doc, "Introduction content guide", 1)
add_heading(doc, "Problem background", 2)
add_para(doc, "Explain that secure application development does not guarantee a secure deployed state. Configuration changes, emergency upgrades, image replacement, port publishing, privilege changes, or dependency replacement can create a difference between the approved baseline and the observed state. Define this difference as application and infrastructure security drift.")
add_para(doc, "Use concrete examples: Java 21 is approved but Java 17 is detected; an internal database is exposed publicly; a non root container is expected but root is configured; the Docker socket is mounted; the running image digest differs; or the dependency version differs from the approved record.")

add_heading(doc, "Research gap", 2)
add_para(doc, "Existing research and commercial tools often focus on one layer such as vulnerable dependencies, SBOM generation, container isolation, configuration validation, exposed services, or runtime enforcement. The project investigates whether a small explainable system can combine a project specific approved baseline with repository and supplied runtime evidence, then preserve the reason and evidence for each mismatch.")

add_heading(doc, "Objectives", 2)
add_bullets(doc, [
    "Store an approved baseline for each application.",
    "Extract supported configuration evidence from user supplied repositories without modifying them.",
    "Accept observed runtime facts through the web interface and REST API.",
    "Detect six classes of baseline drift with deterministic rules.",
    "Prioritize each finding using rule risk, asset criticality, Internet exposure, and known exploitation context.",
    "Preserve evidence, scan history, remediation guidance, and finding status.",
])

add_heading(doc, "Suggested research questions", 2)
add_table(doc, ["ID", "Research question", "Evidence needed"], [
    ["RQ1", "Can deterministic baseline comparison correctly identify the six supported drift categories in controlled cases?", "Execute the 120 cases and compare actual outputs with expected rule IDs."],
    ["RQ2", "Can CodeShield extract usable security evidence from real public repositories?", "Repository scan completion, parsed files, extracted observations, findings, and scan time."],
    ["RQ3", "How does context change finding priority?", "Compare scores while changing only criticality, Internet exposure, or known exploitation."],
    ["RQ4", "Does each finding remain understandable and traceable?", "Review expected state, observed state, evidence source, rule ID, score, and remediation."],
], widths=[0.55, 4.05, 2.65])
add_para(doc, "RQ1 through RQ4 are recommended framing. Only RQ2 has preliminary repository scan observations, and RQ1 has a labeled dataset plus unit tests. Do not write final answers for unexecuted experiments.", italic=True)

add_heading(doc, "Project contributions", 2)
add_bullets(doc, [
    "A simple baseline model that represents approved application configuration.",
    "A user driven repository inspection path for public HTTPS URLs and local directories.",
    "An evidence normalization layer that converts repository or manual observations into the same scan request model.",
    "Six explainable drift rules with stable identifiers and remediation guidance.",
    "A configurable contextual scoring model and a persistent finding workflow.",
    "A deterministic benchmark workbook with controlled compliant and drift cases.",
])
add_para(doc, "Describe these as engineering and evaluation contributions. Do not claim a new vulnerability detection algorithm, a patented invention, or state of the art performance.")

doc.add_page_break()

add_heading(doc, "Relevant work content guide", 1)
add_para(doc, "Organize the related work by research theme instead of summarizing ten papers one by one. Each theme should explain what the literature contributes, its practical limitation for this project, and how CodeShield uses the insight.")

add_heading(doc, "Container threats and isolation", 2)
add_para(doc, "Jarkas et al. provide a broad taxonomy of container vulnerabilities, attacks, and defenses. Zhou et al. propose ARM CCA based container isolation. He et al. demonstrate cross container attack paths created by powerful eBPF capabilities. Together, these papers justify monitoring container users, mounts, privileges, isolation boundaries, and host control paths. Their focus is taxonomy, isolation, or attack demonstration rather than application specific baseline comparison.")

add_heading(doc, "Misconfiguration and network exposure", 2)
add_para(doc, "Choi et al. measure Internet exposed Docker and Kubernetes components and demonstrate the security impact of misconfiguration. Rankovic et al. combine schema validation with configuration version comparison to prevent and diagnose distributed cloud errors. These studies support CodeShield's database exposure and configuration difference rules. CodeShield adds a project specific baseline, but its current parser covers only selected build and container files.")

add_heading(doc, "Dependency reachability and software supply chains", 2)
add_para(doc, "Deng et al. test exploitability through downstream supply chains, while Jia et al. reduce false alarms through vulnerability reachability analysis. Stalnaker et al. document practical SBOM quality, maintenance, and verification problems. These works support the project's decision to treat version differences as evidence that requires context rather than as proof of exploitation. The current prototype does not yet perform call graph analysis, SBOM ingestion, or CVE matching.")

add_heading(doc, "Runtime protection and temporary mitigation", 2)
add_para(doc, "Li et al. generate temporary vulnerability trigger policies, and Kermabon Bobinnec et al. filter malicious system call sequences for unpatched container vulnerabilities. These approaches show how precise runtime evidence can support enforcement. CodeShield currently stops at detection, prioritization, and workflow status; it does not enforce runtime controls.")

add_heading(doc, "Research gap statement", 2)
add_para(doc, "The reviewed studies provide strong methods for individual layers, including taxonomy, reachability, SBOM analysis, exposure measurement, isolation, and runtime filtering. CodeShield focuses on a narrower integration question: whether approved application intent and observed configuration can be represented in a common model and compared through explainable rules. The prototype contributes an integrated workflow rather than replacing the specialized techniques in the literature.")

papers = [
    ["[1]", "Jarkas O et al. 2025  A Container Security Survey  ACM Computing Surveys  DOI 10.1145/3715001", "Threat taxonomy and container control categories"],
    ["[2]", "Deng P et al. 2025  ChainFuzz  USENIX Security 2025", "Downstream exploitability and supply chain context"],
    ["[3]", "Jia Z et al. 2025  Impact Assessment of Third Party Library Vulnerabilities  Computers and Security 157 104546", "Reachability based reduction of dependency false alarms"],
    ["[4]", "Zhou Q et al. 2025  RContainer  NDSS 2025  DOI 10.14722/ndss.2025.240328", "High assurance container isolation"],
    ["[5]", "Li Y et al. 2025  VulShield  NDSS 2025  DOI 10.14722/ndss.2025.240298", "Temporary protection before patch deployment"],
    ["[6]", "Kermabon Bobinnec H et al. 2024  Phoenix  NDSS 2024  DOI 10.14722/ndss.2024.24582", "Runtime syscall sequence filtering"],
    ["[7]", "Choi D et al. 2024  Uncovering Threats in Container Systems  IEEE Access 12", "Measured impact of exposed container components"],
    ["[8]", "Rankovic T et al. 2024  Misconfiguration Prevention and Error Cause Detection  IEEE SISY 2024", "Schema validation and versioned configuration differences"],
    ["[9]", "Stalnaker T et al. 2024  BOMs Away  ICSE 2024  DOI 10.1145/3597503.3623347", "SBOM creation use maintenance and verification challenges"],
    ["[10]", "He Y et al. 2023  Cross Container Attacks  USENIX Security 2023", "eBPF permissions and cross container blast radius"],
]
add_heading(doc, "Literature mapping", 2)
add_table(doc, ["Ref", "Paper", "Use in the report"], papers, widths=[0.55, 4.7, 2.0], font_size=8.6)

doc.add_page_break()

add_heading(doc, "Methodology content guide", 1)
add_heading(doc, "Research and engineering method", 2)
add_para(doc, "Describe the work as design and implementation of a security analysis prototype followed by functional and controlled evaluation. The method begins with an approved baseline, collects or accepts observations, normalizes both into comparable fields, applies deterministic rules, calculates a contextual score, and persists the result.")

add_heading(doc, "Processing flow", 2)
add_numbered(doc, [
    "Register an application baseline containing ownership, criticality, expected Java version, expected database exposure, expected container user, approved image digest, and expected dependency version.",
    "Acquire evidence from a user supplied public HTTPS Git URL, a local path, the manual form, or the REST API.",
    "For remote repositories, validate the URL, reject embedded credentials and private address targets, shallow clone into temporary storage, and remove the temporary copy after inspection.",
    "Search supported files while excluding generated directories and files above the configured size limit.",
    "Extract observations and create a common ScanRequest representation.",
    "Apply six comparison rules and create a finding for every mismatch.",
    "Calculate a score from the rule base risk and context modifiers, map the score to a severity, and store the scan and findings.",
    "Present expected state, observed state, evidence, remediation, status, and scan history in the web interface and REST API.",
])

add_heading(doc, "Repository evidence extraction", 2)
add_table(doc, ["Source", "Extracted observation", "Important behavior"], [
    ["pom.xml", "Java version and Spring Boot parent version", "Uses declared XML properties and parent version patterns."],
    ["build.gradle or build.gradle.kts", "Java and Spring Boot plugin versions", "Uses supported Gradle syntax patterns."],
    ["Dockerfile", "Java base image version and final USER", "If a Dockerfile exists without USER, the observed user is treated as root."],
    ["Compose files", "Port mappings, database exposure, Docker socket mount, and image reference", "Loopback only database mappings are internal. Other matching database port mappings are public."],
    ["Manual or REST input", "All scan fields and contextual flags", "Supports observations that cannot be inferred from repository files."],
], widths=[1.55, 2.8, 2.9])

add_heading(doc, "Implemented drift rules", 2)
add_table(doc, ["Rule ID", "Condition", "Default base risk"], [
    ["CS RUNTIME 001", "Observed Java version differs from the approved version", "30"],
    ["CS NETWORK 001", "An internally approved database is observed as public", "52"],
    ["CS CONTAINER 001", "Observed container user differs from the approved user", "42 for root  25 otherwise"],
    ["CS CONTAINER 002", "Docker socket is mounted inside the application container", "58"],
    ["CS SUPPLY 001", "Observed image digest differs from the approved digest or is unpinned", "38"],
    ["CS DEPENDENCY 001", "Observed dependency version differs from the approved version", "48 with known exploitation  24 otherwise"],
], widths=[1.6, 4.55, 1.1])

add_heading(doc, "Risk scoring", 2)
add_para(doc, "Use the following default formula. Explain that every value is externally configurable in application.yml or through environment variables.")
formula = doc.add_paragraph()
formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
formula.paragraph_format.space_before = Pt(5)
formula.paragraph_format.space_after = Pt(7)
set_font(formula.add_run("risk score = min 100  base rule risk + criticality weight + Internet exposure weight + known exploitation weight"), name="Courier New", size=9.7, bold=True)
add_para(doc, "Criticality weights are Low 0, Medium 8, High 15, and Critical 22. Internet exposure adds 15 and known exploitation adds 18. Scores below 40 are Low, 40 to 64 are Medium, 65 to 84 are High, and 85 to 100 are Critical.")

add_heading(doc, "Data model and software modules", 2)
add_table(doc, ["Component", "Responsibility"], [
    ["ApplicationAsset", "Stores the approved baseline, owner, description, and criticality."],
    ["ScanRun", "Stores the selected application, time, source reference, status, finding count, and maximum score."],
    ["Finding", "Stores rule, explanation, expected and observed states, evidence, remediation, severity, score, and workflow status."],
    ["RepositoryInspectionService", "Validates repository sources, clones public repositories, discovers supported files, and extracts observations."],
    ["DriftDetectionService", "Compares baseline and observed values and emits rule results."],
    ["RiskScoringService", "Adds configurable context weights and maps scores to severity."],
    ["ScanOrchestrator", "Creates scan records, findings, scores, and completion summaries in one workflow."],
    ["Spring MVC and REST controllers", "Provide dashboard, forms, history, finding workflow, and JSON endpoints."],
], widths=[2.05, 5.2])

add_heading(doc, "Safety and reproducibility controls", 2)
add_bullets(doc, [
    "Repository scans do not modify the target repository.",
    "Remote input accepts public HTTPS URLs and rejects URL credentials, fragments, loopback, link local, site local, multicast, and private IPv6 destinations.",
    "Remote repositories use a configurable shallow clone depth and timeout and are deleted from temporary storage after inspection.",
    "Generated directories are ignored and configuration files above the configured size limit are skipped.",
    "A fresh installation starts without seeded assets or findings.",
    "The benchmark dataset is deterministic and includes a generation version and checksum.",
])

doc.add_page_break()

add_heading(doc, "Experimental results content guide", 1)
add_para(doc, "Separate measured results from proposed experiments. Report the evidence below exactly, and do not convert dataset composition into accuracy, precision, recall, or F1 scores until the cases have been executed independently against the application.")

add_heading(doc, "Verified automated test result", 2)
add_table(doc, ["Test area", "Tests", "Verified result"], [
    ["Web and API smoke tests", "4", "Dashboard, asset API, repository scan page, and history page tests passed."],
    ["Drift rule tests", "2", "All six rule categories were detected in a constructed drift case and no finding was produced for a matching case."],
    ["Risk scoring tests", "2", "Context modifiers, score cap, and severity boundaries passed."],
    ["Repository inspection tests", "3", "Evidence extraction, unsupported directory rejection, and non HTTPS remote rejection passed."],
    ["Total", "11", "11 passed  0 failures  0 errors  0 skipped."],
], widths=[2.0, 0.7, 4.55])

add_heading(doc, "Benchmark dataset composition", 2)
add_table(doc, ["Measure", "Value"], [
    ["Dataset version", "CSB 2026 001"],
    ["Total cases", "120"],
    ["Compliant cases", "24"],
    ["Drift cases", "96"],
    ["Expected findings", "181"],
    ["Split", "84 train  18 validation  18 test"],
    ["Highest severity distribution", "None 24  Low 6  Medium 31  High 33  Critical 26"],
    ["Rule frequencies", "Runtime 31  Network 31  Unexpected user 29  Docker socket 30  Image digest 31  Dependency 29"],
], widths=[2.25, 5.0])
add_para(doc, "The dataset provides controlled expected outputs for regression testing. Its labels were generated from the same rule specification used by CodeShield, so it establishes deterministic coverage but does not by itself prove generalization to unknown real world repositories.", italic=True)

add_heading(doc, "Recorded public repository scans", 2)
add_table(doc, ["Repository", "Recorded result", "Interpretation limit"], [
    ["spring projects spring petclinic", "Completed with 5 findings and maximum score 82. Recorded rules covered Java version, database exposure, container user, image digest, and dependency version.", "Findings express differences from the selected stored baseline. They are not a claim that the repository contains five exploitable vulnerabilities."],
    ["ShauryaRawat27 Adaptive LLM", "Completed with 2 findings and maximum score 72. Recorded rules covered container user and an unpinned image reference.", "The result depends on supported files and the chosen baseline. Unsupported configuration is outside the scan scope."],
], widths=[2.05, 3.25, 2.0], font_size=8.8)

add_heading(doc, "Clean installation smoke result", 2)
add_para(doc, "A packaged JAR was started against an empty in memory database. The dashboard, manual scan page, and repository scan page returned HTTP 200, and GET /api/assets returned an empty array. This verifies that the application no longer depends on preloaded demonstration data.")

add_heading(doc, "Experiments still required for stronger claims", 2)
add_bullets(doc, [
    "Execute all 120 benchmark cases through the application and calculate per rule true positives, false positives, false negatives, precision, recall, and F1 score.",
    "Measure repository scan time and extracted file count across a larger and documented set of public repositories.",
    "Create paired scoring cases that change only one context variable to quantify the effect of criticality, exposure, and exploitation evidence.",
    "Ask independent reviewers to judge whether evidence and remediation are understandable, using a defined rubric.",
    "Repeat scans after controlled remediation to verify that findings close consistently.",
])

add_heading(doc, "Result table template for the final report", 2)
add_table(doc, ["Rule", "TP", "FP", "FN", "Precision", "Recall", "F1"], [
    ["Runtime", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
    ["Network", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
    ["Container user", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
    ["Docker socket", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
    ["Image digest", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
    ["Dependency", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
], widths=[1.9, 0.55, 0.55, 0.55, 1.2, 1.2, 1.2], font_size=8.7)

doc.add_page_break()

add_heading(doc, "Conclusions content guide", 1)
add_para(doc, "The conclusion should directly answer the project objective. State that the prototype demonstrates an end to end path from a registered baseline and user supplied evidence to explainable persisted findings. Refer to the passing tests, controlled benchmark coverage, and completed public repository scans as evidence of functional feasibility.")
add_para(doc, "Then state the limits: repository parsing covers selected Java and container files; repository evidence is not equivalent to live runtime observation; the dataset labels are specification derived; the current evaluation does not establish real world precision or recall; and the prototype lacks authentication, autonomous collectors, vulnerability intelligence, SBOM analysis, Kubernetes support, and production persistence.")
add_para(doc, "End with specific future work: execute the benchmark end to end, broaden the repository corpus, add read only authenticated runtime collectors, integrate SBOM and vulnerability sources, measure overhead and accuracy, migrate to PostgreSQL, and add access control. Avoid generic claims that the platform will transform cybersecurity.")

add_heading(doc, "Claims and wording boundaries", 1)
add_table(doc, ["Use this wording", "Avoid this wording", "Reason"], [
    ["Deterministic security drift detection", "AI powered security platform", "The implementation does not use AI or machine learning."],
    ["Repository configuration evidence", "Live runtime state discovered from Git", "A repository scan observes declared configuration, not a running host."],
    ["Public HTTPS repositories and local paths", "Any Git repository", "Private authenticated repositories are not currently supported."],
    ["Six implemented rules", "Complete container security coverage", "The parser and rules cover a defined subset."],
    ["Context adjusted rule score", "Exploit probability", "The score is a transparent priority model, not a calibrated probability."],
    ["120 labeled benchmark cases", "120 real attacks", "The benchmark is synthetic and deterministic."],
    ["Recorded differences from a selected baseline", "Confirmed vulnerabilities in the scanned repositories", "A drift finding is not proof of exploitability."],
    ["Prototype with publication potential", "Novel patented system", "Novelty and patentability require prior art review and stronger evaluation."],
], widths=[2.35, 2.35, 2.55], font_size=8.8)

doc.add_page_break()
add_heading(doc, "Files to attach with the master prompt", 1)
add_bullets(doc, [
    "README.md for the implemented scope, operation, API, rules, and configuration.",
    "CodeShield Literature Review 10 Papers.xlsx for the verified research table and links.",
    "CodeShield Security Drift Benchmark Dataset.xlsx for benchmark cases, labels, scoring, rule catalog, and methodology.",
    "CodeShield Short Project Report.docx only as supporting background. Correct any future architecture statements using this prompt pack.",
    "Source files under src/main/java and src/main/resources when the writing tool can inspect code.",
    "Maven test output or Surefire reports for the final verified test count.",
])

doc.add_page_break()

add_heading(doc, "Master prompt for creating the Review 3 report", 1)
add_para(doc, "Copy the text below and attach the project files listed above. Replace bracketed identity or formatting placeholders before use.")

prompt_paragraphs = [
    "Write a formal Review 3 academic report for the project CodeShield Context Aware Security Drift Analyzer. The required sections are Abstract, Introduction, Relevant Work, Methodology, Experimental Results, and Conclusions. Use clear technical prose suitable for a computer science or cybersecurity review. Use third person or first person plural consistently according to [institutional preference]. Follow [citation style] and [word or page limit].",
    "Treat the supplied source code, README, benchmark workbook, literature review workbook, and verified test output as the evidence base. Do not invent results, performance numbers, user studies, accuracy percentages, citations, components, or deployment capabilities. Distinguish implemented features from proposed architecture and future work.",
    "Describe CodeShield as a deterministic Spring Boot security drift prototype. It stores an approved application baseline, accepts repository configuration or supplied runtime observations, compares expected and observed states through six explainable rules, applies a configurable context adjusted score, and persists evidence linked findings and scan history. It does not use AI or machine learning.",
    "In the Abstract, state the problem, objective, method, Java and Spring Boot implementation, verified results, conclusion, and main limitation in one paragraph of about 180 to 250 words.",
    "In the Introduction, define application and infrastructure security drift, explain why source security can differ from deployed configuration, give concrete examples, identify the gap between single layer scanners and baseline correlation, state the project objectives, list the engineering contributions, and define the current scope.",
    "In Relevant Work, synthesize the ten supplied papers by theme: container threats and isolation; misconfiguration and network exposure; dependency reachability and supply chains; SBOM quality; and runtime protection. Compare their methods and limitations with CodeShield. Do not claim that CodeShield replaces those specialized techniques. End with the narrow research gap addressed by the prototype.",
    "In Methodology, explain the data model, baseline registration, repository source validation, public HTTPS shallow cloning, temporary cleanup, supported file discovery, Maven and Gradle parsing, Dockerfile and Compose extraction, ScanRequest normalization, the six rule conditions, risk scoring, persistence, user interface, REST API, security boundaries, and reproducibility controls. State all repository parser fallback behavior and limitations.",
    "Use the default risk formula: score equals base rule risk plus criticality weight plus 15 for Internet exposure plus 18 for known exploitation, capped at 100. Criticality weights are 0, 8, 15, and 22 for Low, Medium, High, and Critical. Severity thresholds are Medium at 40, High at 65, and Critical at 85. Explain that these values are configuration settings and can be overridden.",
    "In Experimental Results, report that 11 of 11 automated tests passed with zero failures, errors, or skipped tests. Break this into four web and API tests, two drift rule tests, two risk tests, and three repository inspection tests. Report the benchmark composition as 120 deterministic labeled cases, including 24 compliant cases, 96 drift cases, 181 expected findings, and an 84 18 18 train validation test split. Make clear that this describes generated ground truth and does not yet establish measured precision or recall.",
    "Report two recorded public repository scans as preliminary integration observations. The spring projects spring petclinic scan completed with five findings and a maximum score of 82. The ShauryaRawat27 Adaptive LLM scan completed with two findings and a maximum score of 72. Explain that findings are differences from the selected stored baseline and are not claims of exploitable vulnerabilities in those repositories.",
    "Also report that a packaged JAR started against an empty database and returned HTTP 200 for the dashboard, manual scan page, and repository scan page, while the asset API returned an empty array. This verifies that the application does not require preloaded sample data.",
    "Do not calculate precision, recall, F1, false positive rate, performance improvement, or statistical significance unless separate executed results are supplied. Instead, state the exact experiments required to calculate them and use clearly labeled TBD fields where necessary.",
    "In Conclusions, state what the current evidence supports, list the implementation limits, and propose specific next work: execute the 120 case benchmark end to end, expand the repository corpus, add read only authenticated runtime collectors, integrate SBOM and vulnerability intelligence, measure scan overhead and accuracy, add access control, and use PostgreSQL for production deployment.",
    "Keep the tone factual. Avoid promotional language, claims of patentability, statements that the project is state of the art, and claims of continuous monitoring. Use tables only where they improve comparison. Include a complete reference list for all ten supplied papers using their verified DOI or official conference URLs.",
]
for paragraph_text in prompt_paragraphs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    set_font(p.add_run(paragraph_text), name="Courier New", size=9.2, color=TEXT)

add_heading(doc, "Verified references", 1)
references = [
    ("[1] Jarkas O Ko R Dong N and Mahmud R 2025 A Container Security Survey Exploits Attacks and Defenses ACM Computing Surveys 57 7 Article 170", "https://doi.org/10.1145/3715001"),
    ("[2] Deng P Zhang L Meng Y et al. 2025 ChainFuzz Exploiting Upstream Vulnerabilities in Open Source Supply Chains 34th USENIX Security Symposium", "https://www.usenix.org/conference/usenixsecurity25/presentation/deng"),
    ("[3] Jia Z Yang C Feng P et al. 2025 Impact Assessment of Third Party Library Vulnerabilities through Vulnerability Reachability Analysis Computers and Security 157 104546", "https://doi.org/10.1016/j.cose.2025.104546"),
    ("[4] Zhou Q Cao W Jia X et al. 2025 RContainer A Secure Container Architecture through Extending ARM CCA Hardware Primitives NDSS 2025", "https://doi.org/10.14722/ndss.2025.240328"),
    ("[5] Li Y Zhang C Zhu J et al. 2025 VulShield Protecting Vulnerable Code Before Deploying Patches NDSS 2025", "https://doi.org/10.14722/ndss.2025.240298"),
    ("[6] Kermabon Bobinnec H Jarraya Y Wang L et al. 2024 Phoenix Surviving Unpatched Vulnerabilities via Accurate and Efficient Filtering of Syscall Sequences NDSS 2024", "https://doi.org/10.14722/ndss.2024.24582"),
    ("[7] Choi D Seo H Kim K et al. 2024 Uncovering Threats in Container Systems A Study on Misconfigured Container Components in the Wild IEEE Access 12 192931 to 192945", "https://doi.org/10.1109/ACCESS.2024.3514751"),
    ("[8] Rankovic T Siljic F Tomic J Sladic G and Simic M 2024 Misconfiguration Prevention and Error Cause Detection for Distributed Cloud Applications IEEE SISY 2024", "https://doi.org/10.1109/SISY62279.2024.10737513"),
    ("[9] Stalnaker T Wintersgill N Chaparro O et al. 2024 BOMs Away Inside the Minds of Stakeholders ICSE 2024", "https://doi.org/10.1145/3597503.3623347"),
    ("[10] He Y Guo R Xing Y et al. 2023 Cross Container Attacks The Bewildered eBPF on Clouds 32nd USENIX Security Symposium", "https://www.usenix.org/conference/usenixsecurity23/presentation/he"),
]
for citation, url in references:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.05
    set_font(p.add_run(citation + "  "), size=9.2)
    add_hyperlink(p, "Verified source", url)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
