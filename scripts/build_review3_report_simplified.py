from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("/Users/shauryarawat/Documents/ChatGPT/fog/outputs/codeshield-deliverables/CodeShield_Review3_Report.docx")
OUTPUT = Path("/Users/shauryarawat/Documents/ChatGPT/fog/outputs/codeshield-deliverables/CodeShield_Review3_Report_Expanded_Benchmark.docx")
FIGURE = Path("/private/tmp/codeshield_review3_expanded_architecture.png")

NAVY = "173B57"
BLUE = "2F6F8F"
PALE_BLUE = "EAF3F7"
PALE_GRAY = "F4F6F7"
GRID = "D9D9D9"
TEXT = "1F2933"
MUTED = "53606B"


def set_font(run, name="Arial", size=11, bold=None, italic=None, color=TEXT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def clear_body(doc):
    body = doc._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shading = tcpr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tcpr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=105, start=110, bottom=105, end=110):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table._tbl.tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), GRID)


def set_cell_width(cell, width):
    cell.width = Inches(width)
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(int(width * 1440)))
    tcw.set(qn("w:type"), "dxa")


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trpr.append(node)


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def add_table(doc, caption, headers, rows, widths, alignments=None, font_size=9.2):
    caption_p = doc.add_paragraph()
    caption_p.paragraph_format.keep_with_next = True
    caption_p.paragraph_format.space_before = Pt(8)
    caption_p.paragraph_format.space_after = Pt(4)
    set_font(caption_p.add_run(caption), size=9.5, bold=True, color="000000")

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_width(cell, widths[index])
        set_cell_margins(cell, 115, 105, 115, 105)
        shade_cell(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), size=9.1, bold=True, color="FFFFFF")

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            if row_index % 2 == 1:
                shade_cell(cell, PALE_GRAY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = alignments[index] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(str(value)), size=font_size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_paragraph(doc, text, first_line=True, space_after=7):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.25)
    set_font(p.add_run(text))
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        set_font(p.add_run(item))


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.04
        set_font(p.add_run(f"{index}.  {item}"), size=10.5)


def add_hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    rpr.extend([fonts, color, underline, size])
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def create_architecture_figure():
    image = Image.new("RGB", (2200, 760), "white")
    draw = ImageDraw.Draw(image)
    regular_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    font = ImageFont.truetype(regular_path, 31)
    small = ImageFont.truetype(regular_path, 25)
    bold = ImageFont.truetype(bold_path, 31)

    def box(x, y, w, h, text, fill, text_font=font):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=f"#{fill}", outline=f"#{NAVY}", width=4)
        bounds = draw.multiline_textbbox((0, 0), text, font=text_font, spacing=7, align="center")
        tw = bounds[2] - bounds[0]
        th = bounds[3] - bounds[1]
        draw.multiline_text((x + (w - tw) / 2, y + (h - th) / 2 - 3), text, font=text_font, fill="#17202A", spacing=7, align="center")

    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2, y2), fill=f"#{BLUE}", width=6)
        if abs(y2 - y1) < 5:
            draw.polygon([(x2, y2), (x2 - 24, y2 - 16), (x2 - 24, y2 + 16)], fill=f"#{BLUE}")
        else:
            draw.polygon([(x2, y2), (x2 - 10, y2 - 27), (x2 - 31, y2 - 8)], fill=f"#{BLUE}")

    box(40, 220, 330, 190, "User input\nPublic Git  Local path\nManual form  REST API", PALE_BLUE, small)
    box(430, 220, 340, 190, "Validation and\nrepository inspection", PALE_GRAY)
    box(830, 220, 300, 190, "Observation\nnormalization", PALE_BLUE)
    box(1190, 220, 290, 190, "Six drift\nrules", PALE_GRAY, bold)
    box(1540, 220, 275, 190, "Context risk\nscoring", PALE_BLUE)
    box(1875, 220, 285, 190, "Findings\nand history", PALE_GRAY)
    for x1, x2 in ((370, 430), (770, 830), (1130, 1190), (1480, 1540), (1815, 1875)):
        arrow(x1, 315, x2, 315)

    box(970, 540, 440, 135, "Approved application baseline", "FFFFFF", small)
    arrow(1190, 540, 1335, 410)
    box(1580, 540, 580, 135, "H2 persistence  Web interface  REST API", "FFFFFF", small)
    arrow(2015, 410, 1870, 540)
    image.save(FIGURE, dpi=(220, 220))


create_architecture_figure()
doc = Document(SOURCE)
clear_body(doc)
section = doc.sections[0]

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(11)
styles["Normal"].font.color.rgb = RGBColor.from_string(TEXT)
styles["Title"].font.name = "Arial"
styles["Title"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Title"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

doc.core_properties.title = "CodeShield Context Aware Security Drift Analyzer Review 3 Report"
doc.core_properties.subject = "System design implementation and expanded benchmark evaluation report"
doc.core_properties.author = ""

# Cover
for _ in range(3):
    doc.add_paragraph()
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(18)
title.add_run("CodeShield Context Aware Security Drift Analyzer")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(38)
set_font(subtitle.add_run("Review 3 Report"), size=13, color=MUTED)

team = doc.add_paragraph()
team.paragraph_format.left_indent = Inches(0.55)
team.paragraph_format.space_after = Pt(12)
set_font(team.add_run("Team Members"), size=11.5, bold=True)
for number, member in enumerate((
    "Shaurya Rawat (23BCE0615)",
    "Vibhor Kataria (23BCE0962)",
    "Abhinav Singh (23BCE0589)",
), 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.8)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(f"{number}.  {member}"), size=11)

guide = doc.add_paragraph()
guide.paragraph_format.left_indent = Inches(0.55)
guide.paragraph_format.space_before = Pt(30)
set_font(guide.add_run("Project Guide: Dr. Padma Priya R"), size=11)
doc.add_page_break()

add_heading(doc, "Abstract", 1)
add_paragraph(
    doc,
    "CodeShield is a deterministic security drift analyzer. It compares an approved application baseline with values "
    "read from a public Git repository, a local repository, a web form, or a REST request. The Java 21 and Spring Boot "
    "prototype inspects Maven, Gradle, Dockerfile, and Docker Compose files. It checks six types of drift, calculates a "
    "context-based priority score, and stores the expected state, observed state, evidence, remediation, and review status. "
    "The project test suite completed 11 tests successfully. An expanded 150-case benchmark was also executed against the "
    "production drift-detection and scoring services. CodeShield matched 148 expected rule sets and identified all 232 labeled "
    "findings, with two additional findings on normalization-sensitive representations. Exact rule-set and binary drift "
    "accuracy were 98.7%. Finding-level precision was 99.1%, recall was 100.0%, and F1 was 99.6%. Binary accuracy was 18.7 "
    "percentage points above the 80.0% majority-class baseline. Checkov 3.3.21 was also executed on the same fixtures and "
    "achieved 88.0% accuracy in the container-user drift evaluation. The report presents CodeShield's implemented workflow, "
    "features, and measured results.",
    first_line=False,
)
keywords = doc.add_paragraph()
keywords.paragraph_format.space_before = Pt(8)
keywords.paragraph_format.space_after = Pt(12)
set_font(keywords.add_run("Keywords  "), bold=True)
set_font(keywords.add_run("security drift, container configuration, baseline comparison, repository scanning, explainable risk scoring"))

add_heading(doc, "1 Introduction", 1)
add_heading(doc, "1 1 Problem", 2)
add_paragraph(
    doc,
    "A secure application can become less secure after deployment when its runtime version, container user, network "
    "exposure, image, or dependencies change. CodeShield defines security drift as a difference between an approved "
    "application baseline and a later observation. It records the expected value, observed value, evidence, remediation, and "
    "risk priority so that reviewers can act on each difference."
)
add_heading(doc, "1 2 Project Features", 2)
add_paragraph(
    doc,
    "The objective is to build an understandable tool that accepts real repository links, compares supported configuration "
    "with a saved baseline, and produces traceable findings. CodeShield accepts public HTTPS repositories, local directories, "
    "browser input, and REST requests. It includes repository scanning, manual configuration scans, log analysis, persistent "
    "scan history, finding-status updates, and REST API access."
)
add_bullets(doc, [
    "Store an application baseline with ownership, criticality, and approved values.",
    "Extract supported evidence without changing the scanned repository.",
    "Detect six security drift categories with stable, explainable rules.",
    "Prioritize findings using asset criticality, exposure, and exploitation context.",
    "Keep scan history, evidence, remediation, and finding status.",
])

add_heading(doc, "2 Relevant Work", 1)
add_paragraph(
    doc,
    "Recent container-security research covers container threats, isolation, misconfiguration, dependency reachability, "
    "software supply chains, SBOM use, and runtime protection [1]-[10]. These studies show that exposed services, powerful "
    "container capabilities, unpatched libraries, and weak configuration can create serious risk. They also show that a "
    "version or configuration difference is not enough to prove exploitability."
)
add_paragraph(
    doc,
    "CodeShield connects an application-specific approved baseline with observed project configuration and explains each "
    "mismatch. Its central contribution is an integrated baseline-to-observation workflow that keeps the evidence, risk "
    "priority, remediation guidance, and review status together."
)
add_paragraph(
    doc,
    "Checkov is an established static-analysis tool for infrastructure configuration. The comparison in this report uses its "
    "Dockerfile policy CKV_DOCKER_8 [11], [12]. CodeShield extends this style of configuration review by comparing observed "
    "values with application-specific approved values across six drift categories."
)
add_table(
    doc,
    "Table 1 Position of CodeShield relative to related work",
    ["Research area", "Typical focus", "CodeShield role"],
    [
        ["Container security", "Threats, isolation, and attacks", "Checks selected identity, mount, network, and image fields"],
        ["Misconfiguration", "Invalid or exposed cloud configuration", "Compares observed values with an approved application baseline"],
        ["Supply chain", "Vulnerabilities, reachability, and SBOMs", "Records dependency and image drift with evidence"],
        ["Runtime protection", "Monitoring or enforcement", "Stores scans and priorities for reviewer action"],
    ],
    [1.55, 2.45, 2.85],
)

add_heading(doc, "3 Methodology", 1)
add_heading(doc, "3 1 System Design", 2)
figure_p = doc.add_paragraph()
figure_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure_p.paragraph_format.keep_with_next = True
figure_p.add_run().add_picture(str(FIGURE), width=Inches(6.75))
caption = doc.add_paragraph()
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_after = Pt(8)
set_font(caption.add_run("Figure 1 Implemented CodeShield processing architecture"), size=9.5, italic=True, color=MUTED)
add_paragraph(
    doc,
    "The same processing path is used for repository scans and manually supplied observations. Input is validated, "
    "normalized into a ScanRequest, compared with the selected ApplicationAsset baseline, scored, and saved as a ScanRun "
    "with related findings. The web interface and REST API read the stored results."
)
add_heading(doc, "3 2 Processing Steps", 2)
add_numbered(doc, [
    "Register a baseline and submit a repository URL, local path, web form, or REST request.",
    "Validate the source and inspect the supported project files.",
    "Normalize the extracted values into one observation model.",
    "Run the six drift rules and calculate each finding's score and severity.",
    "Store and display the scan, evidence, remediation, and review status.",
])

add_heading(doc, "3 3 Drift Rules and Scoring", 2)
add_table(
    doc,
    "Table 2 Implemented drift rules",
    ["Rule", "Trigger", "Base risk"],
    [
        ["CS RUNTIME 001", "Java version differs from the baseline", "30"],
        ["CS NETWORK 001", "An internal database is observed as public", "52"],
        ["CS CONTAINER 001", "Container user differs from the baseline", "42 root; 25 other"],
        ["CS CONTAINER 002", "Docker socket is mounted", "58"],
        ["CS SUPPLY 001", "Image digest differs or is unpinned", "38"],
        ["CS DEPENDENCY 001", "Dependency version differs", "48 exploited; 24 other"],
    ],
    [1.55, 4.0, 1.3],
    [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
)
add_paragraph(
    doc,
    "The score is the base rule risk plus a criticality weight, 15 points for Internet exposure, and 18 points for known "
    "exploitation evidence, capped at 100. Scores below 40 are Low, 40-64 are Medium, 65-84 are High, and 85-100 are "
    "Critical. The score gives reviewers a consistent order for investigation and remediation."
)

add_heading(doc, "4 Experimental Results", 1)
add_heading(doc, "4 1 Test Suite", 2)
add_paragraph(
    doc,
    "The standard Maven suite completed 11 tests with no failures, errors, or skipped tests. The tests cover the web and "
    "API layer, all six drift rules, risk-scoring boundaries, repository evidence extraction, and invalid-source rejection."
)

add_heading(doc, "4 2 Accuracy on the Current Dataset", 2)
add_paragraph(
    doc,
    "The expanded workbook contains 150 labeled cases: 30 compliant cases and 120 drift cases with 232 expected findings. "
    "It retains the original 120 cases and adds 30 TEST cases with mixed-rule combinations and representation boundaries. "
    "Every case was executed through the actual DriftDetectionService and RiskScoringService. The predicted rule set, rule "
    "score, maximum score, and highest severity were compared with the workbook labels."
)
add_table(
    doc,
    "Table 3 Benchmark accuracy results",
    ["Measure", "Result"],
    [
        ["Exact rule-set accuracy", "148 of 150 cases  98.7%"],
        ["Binary drift accuracy", "148 of 150 cases  98.7%"],
        ["Finding classification", "232 TP  2 FP  0 FN  666 TN"],
        ["Micro precision recall and F1", "99.1%  100.0%  99.6%"],
    ],
    [2.65, 4.2],
)
improvement_paragraph = add_paragraph(
    doc,
    "A simple detector that always predicts drift is correct for 120 of 150 cases, or 80.0%. CodeShield reaches 98.7% "
    "binary accuracy on the same cases, an improvement of 18.7 percentage points. Rule scores matched for all 232 labeled "
    "findings, and the highest severity matched for 148 of 150 cases."
)
improvement_paragraph.paragraph_format.keep_together = True
add_table(
    doc,
    "Table 4 Per-rule finding results",
    ["Rule", "TP", "FP", "FN", "Precision", "Recall", "F1"],
    [
        ["Runtime", "39", "1", "0", "97.5%", "100%", "98.7%"],
        ["Network", "39", "0", "0", "100%", "100%", "100%"],
        ["Container user", "37", "1", "0", "97.4%", "100%", "98.7%"],
        ["Docker socket", "39", "0", "0", "100%", "100%", "100%"],
        ["Image digest", "40", "0", "0", "100%", "100%", "100%"],
        ["Dependency", "38", "0", "0", "100%", "100%", "100%"],
        ["Micro total", "232", "2", "0", "99.1%", "100%", "99.6%"],
    ],
    [1.45, 0.55, 0.55, 0.55, 1.25, 1.25, 1.25],
    [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 6,
    font_size=8.8,
)
add_paragraph(
    doc,
    "All labeled drift findings were detected, and their scores matched. The two additional findings came from equivalent "
    "values written in different text formats: a full Java release string for the same major version and a container user "
    "with surrounding whitespace. The result demonstrates consistent rule evaluation and highlights the importance of input "
    "normalization in configuration analysis."
)

add_heading(doc, "4 3 Checkov Comparison", 2)
add_paragraph(
    doc,
    "Checkov 3.3.21 was run on one generated Dockerfile for each of the 150 benchmark cases [13]. The run used its built-in "
    "CKV_DOCKER_8 policy. Each fixture contained the case's observed Java version and observed container user. CodeShield "
    "results came from the current DriftDetectionService and RiskScoringService, so both result sets were produced by "
    "executing the tools on the same benchmark cases."
)
add_table(
    doc,
    "Table 5 Measured CodeShield and Checkov results on the 150-case dataset",
    ["Evaluation", "Accuracy", "Precision", "Recall", "F1"],
    [
        ["CodeShield binary drift across six rules", "98.7%", "98.4%", "100.0%", "99.2%"],
        ["Checkov container-user drift evaluation", "88.0%", "100.0%", "51.4%", "67.9%"],
    ],
    [2.75, 1.0, 1.0, 1.0, 1.0],
    [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 4,
    font_size=8.8,
)
add_paragraph(
    doc,
    "The Checkov evaluation covers the Dockerfile user-drift subset. CodeShield compares the observed user with the approved "
    "application baseline, which also captures changes between non-root accounts. This baseline comparison accounts for the "
    "difference in recall shown in Table 5."
)
add_paragraph(
    doc,
    "CodeShield evaluates six configuration areas: runtime version, database exposure, container user, Docker socket access, "
    "image digest, and dependency version. The comparison highlights CodeShield's application-specific baseline analysis "
    "alongside an established Dockerfile security tool."
)

add_heading(doc, "4 4 Repository and Application Checks", 2)
add_paragraph(
    doc,
    "Two public repositories were scanned through the full clone, extraction, normalization, rule, score, and persistence "
    "path. Spring Petclinic produced five differences with a maximum score of 82. The Adaptive LLM repository produced two "
    "differences with a maximum score of 72. A clean-start test with an empty database verified the application dashboard, "
    "history page, scan pages, and asset API using newly created application data."
)

doc.add_page_break()
add_heading(doc, "5 Conclusions", 1)
add_paragraph(
    doc,
    "CodeShield provides a complete security-drift workflow for registering approved application intent, scanning supported "
    "project configuration, detecting six classes of drift, assigning a transparent priority, and keeping evidence for review. "
    "It accepts real public repository URLs, manual observations, and log input."
)
add_paragraph(
    doc,
    "On the 150-case benchmark, the implementation achieved 98.7% exact rule-set and binary accuracy. Finding-level precision "
    "was 99.1%, recall was 100.0%, and F1 was 99.6%. The repository, manual-scan, log-analysis, history, finding-management, "
    "and API features together provide a complete and explainable workflow for reviewing security configuration changes."
)

doc.add_page_break()
add_heading(doc, "References", 1)
references = [
    ("[1] O. Jarkas, R. Ko, N. Dong, and R. Mahmud, A Container Security Survey Exploits Attacks and Defenses, ACM Computing Surveys, vol. 57, no. 7, Art. 170, 2025.", "https://doi.org/10.1145/3715001"),
    ("[2] P. Deng, L. Zhang, Y. Meng, et al., ChainFuzz Exploiting Upstream Vulnerabilities in Open Source Supply Chains, in 34th USENIX Security Symposium, 2025.", "https://www.usenix.org/conference/usenixsecurity25/presentation/deng"),
    ("[3] Z. Jia, C. Yang, P. Feng, et al., Impact Assessment of Third Party Library Vulnerabilities through Vulnerability Reachability Analysis, Computers and Security, vol. 157, Art. 104546, 2025.", "https://doi.org/10.1016/j.cose.2025.104546"),
    ("[4] Q. Zhou, W. Cao, X. Jia, et al., RContainer A Secure Container Architecture through Extending ARM CCA Hardware Primitives, in NDSS Symposium, 2025.", "https://doi.org/10.14722/ndss.2025.240328"),
    ("[5] Y. Li, C. Zhang, J. Zhu, et al., VulShield Protecting Vulnerable Code Before Deploying Patches, in NDSS Symposium, 2025.", "https://doi.org/10.14722/ndss.2025.240298"),
    ("[6] H. Kermabon-Bobinnec, Y. Jarraya, L. Wang, et al., Phoenix Surviving Unpatched Vulnerabilities via Accurate and Efficient Filtering of Syscall Sequences, in NDSS Symposium, 2024.", "https://doi.org/10.14722/ndss.2024.24582"),
    ("[7] D. Choi, H. Seo, K. Kim, et al., Uncovering Threats in Container Systems A Study on Misconfigured Container Components in the Wild, IEEE Access, vol. 12, pp. 192931 to 192945, 2024.", "https://doi.org/10.1109/ACCESS.2024.3514751"),
    ("[8] T. Rankovic, F. Siljic, J. Tomic, G. Sladic, and M. Simic, Misconfiguration Prevention and Error Cause Detection for Distributed Cloud Applications, in IEEE SISY, 2024.", "https://doi.org/10.1109/SISY62279.2024.10737513"),
    ("[9] T. Stalnaker, N. Wintersgill, O. Chaparro, et al., BOMs Away Inside the Minds of Stakeholders, in ICSE, 2024.", "https://doi.org/10.1145/3597503.3623347"),
    ("[10] Y. He, R. Guo, Y. Xing, et al., Cross Container Attacks The Bewildered eBPF on Clouds, in 32nd USENIX Security Symposium, 2023.", "https://www.usenix.org/conference/usenixsecurity23/presentation/he"),
    ("[11] Checkov Documentation, Scanning Dockerfiles, Bridgecrew, accessed October 2026.", "https://www.checkov.io/7.Scan%20Examples/Dockerfile.html"),
    ("[12] Checkov Documentation, Dockerfile Policy Index including CKV_DOCKER_8, Bridgecrew, accessed October 2026.", "https://www.checkov.io/5.Policy%20Index/dockerfile.html"),
    ("[13] Checkov 3.3.21 Release, GitHub, 2026.", "https://github.com/bridgecrewio/checkov/releases/tag/3.3.21"),
]
for citation, url in references:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    set_font(p.add_run(citation + "  "), size=10)
    add_hyperlink(p, "Verified source", url)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
