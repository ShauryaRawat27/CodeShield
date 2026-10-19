from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("/Users/shauryarawat/Documents/ChatGPT/fog/outputs/codeshield-deliverables/CodeShield_Review3_Report.docx")
FIGURE = Path("/private/tmp/codeshield_review3_architecture.png")

NAVY = "173B57"
BLUE = "2F6F8F"
PALE_BLUE = "EAF3F7"
PALE_GRAY = "F4F6F7"
GRID = "D9D9D9"
TEXT = "1F2933"
MUTED = "53606B"


def set_font(run, name="Arial", size=11, bold=None, italic=None, color=TEXT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def add_hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    rpr.extend([fonts, color, underline, size])
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def remove_paragraph_border(style):
    ppr = style.element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shading = tcpr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tcpr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=105, start=110, bottom=105, end=110):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), GRID)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trpr.append(node)


def prevent_row_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def set_cell_width(cell, width):
    cell.width = Inches(width)
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(int(width * 1440)))
    tcw.set(qn("w:type"), "dxa")


def add_table(doc, caption, headers, rows, widths, alignments=None, font_size=9.3):
    caption_p = doc.add_paragraph()
    caption_p.paragraph_format.keep_with_next = True
    caption_p.paragraph_format.space_before = Pt(8)
    caption_p.paragraph_format.space_after = Pt(4)
    set_font(caption_p.add_run(caption), size=9.5, bold=True, color="000000")

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_width(cell, widths[index])
        set_cell_margins(cell, 115, 105, 115, 105)
        shade_cell(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), size=9.2, bold=True, color="FFFFFF")

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            if row_index % 2 == 1:
                shade_cell(cell, PALE_GRAY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = (alignments[index] if alignments else WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(str(value)), size=font_size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_paragraph(doc, text, italic=False, first_line=True, space_after=7):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.25)
    set_font(p.add_run(text), italic=italic)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.1
        set_font(p.add_run(item))


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.1
        set_font(p.add_run(f"{index}.  {item}"))


def create_architecture_figure():
    image = Image.new("RGB", (2200, 960), "white")
    draw = ImageDraw.Draw(image)
    regular_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    font = ImageFont.truetype(regular_path, 32)
    small = ImageFont.truetype(regular_path, 27)
    bold = ImageFont.truetype(bold_path, 32)

    def box(x, y, w, h, text, fill, text_font=font):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=f"#{fill}", outline=f"#{NAVY}", width=4)
        bounds = draw.multiline_textbbox((0, 0), text, font=text_font, spacing=7, align="center")
        tw = bounds[2] - bounds[0]
        th = bounds[3] - bounds[1]
        draw.multiline_text((x + (w - tw) / 2, y + (h - th) / 2 - 3), text, font=text_font, fill="#17202A", spacing=7, align="center")

    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2, y2), fill=f"#{BLUE}", width=6)
        angle = 16
        if abs(y2 - y1) < 5:
            draw.polygon([(x2, y2), (x2 - 24, y2 - angle), (x2 - 24, y2 + angle)], fill=f"#{BLUE}")
        else:
            draw.polygon([(x2, y2), (x2 - 10, y2 - 27), (x2 - 31, y2 - 8)], fill=f"#{BLUE}")

    box(40, 310, 330, 210, "User input\nPublic Git  Local path\nManual form  REST API", PALE_BLUE, small)
    box(430, 310, 340, 210, "Validation and\nrepository inspection", PALE_GRAY)
    box(830, 310, 300, 210, "Observation\nnormalization", PALE_BLUE)
    box(1190, 310, 290, 210, "Six drift\nrules", PALE_GRAY, bold)
    box(1540, 310, 275, 210, "Context risk\nscoring", PALE_BLUE)
    box(1875, 310, 285, 210, "Findings\nand history", PALE_GRAY)
    for x1, x2 in ((370, 430), (770, 830), (1130, 1190), (1480, 1540), (1815, 1875)):
        arrow(x1, 415, x2, 415)

    box(970, 680, 440, 150, "Approved application baseline", "FFFFFF", small)
    arrow(1190, 680, 1335, 520)
    box(1580, 680, 580, 150, "H2 persistence  Web interface  REST API", "FFFFFF", small)
    arrow(2015, 520, 1870, 680)

    image.save(FIGURE, dpi=(220, 220))


create_architecture_figure()

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)
section.different_first_page_header_footer = True

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(11)
styles["Normal"].font.color.rgb = RGBColor.from_string(TEXT)
styles["Title"].font.name = "Arial"
styles["Title"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Title"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
remove_paragraph_border(styles["Title"])
for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

doc.core_properties.title = "CodeShield Context Aware Security Drift Analyzer Review 3 Report"
doc.core_properties.subject = "Design implementation and controlled evaluation of CodeShield"
doc.core_properties.author = ""
doc.core_properties.keywords = "CodeShield, security drift, container security, repository inspection, risk scoring"

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer_p.add_run("CodeShield Review 3 Report    "), size=9, color=MUTED)
add_page_number(footer_p)

# Cover
for _ in range(3):
    doc.add_paragraph()
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(12)
title.add_run("CodeShield Context Aware Security Drift Analyzer")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(34)
set_font(subtitle.add_run("Review 3 Report on Design Implementation and Controlled Evaluation"), size=13, color=MUTED)

for label in (
    "Student Name  [Enter name]",
    "Roll Number  [Enter roll number]",
    "Institution  [Enter institution]",
    "Project Guide  [Enter guide name]",
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run(label), size=11)

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
date_p.paragraph_format.space_before = Pt(24)
set_font(date_p.add_run("October 2026"), size=11, color=MUTED)

doc.add_page_break()

add_heading(doc, "Abstract", 1)
abstract = (
    "Application security controls can drift after development because deployment configuration, container settings, "
    "runtime versions, and dependencies continue to change. This report presents CodeShield, a deterministic security "
    "drift analyzer that compares an approved application baseline with configuration evidence extracted from a public "
    "HTTPS Git repository, a local repository, or manually supplied runtime observations. The prototype is implemented "
    "with Java 21 and Spring Boot 3.3.6. It normalizes observations from Maven, Gradle, Dockerfile, and Docker Compose "
    "files, applies six explainable comparison rules, and assigns a configurable score based on rule risk, asset "
    "criticality, Internet exposure, and known exploitation evidence. Findings retain the expected state, observed state, "
    "evidence, remediation guidance, severity, and workflow status. Eleven automated tests passed without failure, error, "
    "or omission. A deterministic benchmark contains 120 labeled cases and 181 expected findings, while two preliminary "
    "public repository scans completed with five and two baseline differences respectively. A clean-start smoke test also "
    "confirmed that the packaged application operates without preloaded demonstration data. These results establish the "
    "functional feasibility of an evidence-linked baseline comparison workflow. They do not establish real-world "
    "precision, recall, exploitability, or continuous runtime monitoring, which require broader independent evaluation and "
    "additional collection mechanisms."
)
add_paragraph(doc, abstract, first_line=False)
keywords = doc.add_paragraph()
keywords.paragraph_format.space_before = Pt(10)
keywords.paragraph_format.space_after = Pt(14)
set_font(keywords.add_run("Keywords  "), bold=True)
set_font(keywords.add_run("security drift, container configuration, baseline comparison, repository inspection, explainable risk scoring"))

add_heading(doc, "Report Scope", 2)
add_paragraph(
    doc,
    "This report documents the implemented CodeShield prototype and the evidence currently available for Review 3. "
    "Repository scans observe declared configuration rather than live host state, and their findings identify differences "
    "from a selected baseline rather than confirmed exploitable vulnerabilities. Proposed runtime collectors, private "
    "repository authentication, vulnerability intelligence, SBOM analysis, Kubernetes support, access control, and "
    "production persistence are therefore discussed only as future work.",
    first_line=False,
)

doc.add_page_break()

add_heading(doc, "1 Introduction", 1)
add_heading(doc, "1 1 Security Drift Problem", 2)
add_paragraph(
    doc,
    "Secure development practices establish an intended application state, but they do not guarantee that later "
    "configuration will remain identical to that state. Emergency upgrades, image replacement, environment-specific "
    "Compose files, port publishing, permission changes, and dependency replacement can alter the deployed security "
    "posture. In this report, security drift means a measurable difference between an approved application baseline and "
    "subsequently observed repository configuration or supplied runtime evidence."
)
add_paragraph(
    doc,
    "The consequences depend on the field that changes. A Java 21 baseline may be compared with Java 17 configuration; "
    "an internal database may become publicly mapped; a non-root container expectation may be replaced by root; the "
    "Docker socket may be mounted into an application container; a previously approved image digest may be replaced by "
    "an unpinned tag; or a dependency version may change without a corresponding baseline update. Each difference is "
    "useful evidence, but a difference alone does not prove exploitability. It must be interpreted in the context of the "
    "application, its exposure, and the reason for the approved state."
)
add_paragraph(
    doc,
    "Existing security tools often examine one layer at a time. Dependency scanners search for vulnerable packages, "
    "container scanners inspect images, configuration validators apply schemas, and runtime defenses observe or block "
    "behavior. These functions remain important, but they do not always express whether a specific application has moved "
    "away from its own approved configuration. CodeShield addresses this narrower integration problem by storing intent, "
    "normalizing observations, comparing fields through stable rules, and preserving the evidence behind each result."
)

add_heading(doc, "1 2 Objectives and Research Questions", 2)
add_paragraph(
    doc,
    "The project objective is to design and implement an explainable prototype that detects selected classes of security "
    "drift and provides enough evidence for a user to understand and review each mismatch. The implementation pursues the "
    "following engineering objectives.",
)
add_bullets(doc, [
    "Store an approved application baseline with ownership and criticality information.",
    "Extract supported configuration evidence from user-supplied public repositories and local directories without modifying the source.",
    "Accept observations through a browser interface and REST API when repository files cannot represent the required field.",
    "Detect six drift categories with deterministic rules and stable identifiers.",
    "Prioritize findings through configurable rule and context weights.",
    "Preserve expected state, observed state, evidence, remediation, history, and finding status.",
])
add_paragraph(
    doc,
    "The current evaluation is organized around four questions: whether deterministic baseline comparison correctly "
    "represents the six supported categories in controlled cases; whether the repository inspector extracts usable evidence "
    "from real public projects; how criticality, exposure, and exploitation context affect finding priority; and whether "
    "each result remains traceable to an expected value, observed value, rule, and evidence source. The completed work "
    "provides unit and integration evidence for these questions, but it does not yet provide an independent accuracy study."
)

add_heading(doc, "1 3 Contributions and Scope", 2)
add_paragraph(
    doc,
    "CodeShield contributes an integrated engineering workflow rather than a new vulnerability-discovery algorithm. The "
    "workflow combines an application-specific baseline, repository or manually supplied evidence, a common scan request "
    "model, six stable comparison rules, contextual scoring, and persistent finding management. A deterministic benchmark "
    "workbook accompanies the implementation so that the supported rule combinations can be reviewed and executed "
    "repeatedly."
)
add_paragraph(
    doc,
    "The current scope is deliberately limited. Public HTTPS clone URLs and readable local paths are supported, while "
    "authenticated private repositories are not. Parsing covers selected Maven, Gradle, Dockerfile, and Compose patterns. "
    "The system uses an embedded H2 database and does not include user authentication, Kubernetes or Terraform analysis, "
    "live CVE feeds, SBOM ingestion, autonomous host agents, or automatic remediation. These boundaries keep the Review 3 "
    "claims aligned with the implemented code."
)

add_heading(doc, "2 Relevant Work", 1)
add_heading(doc, "2 1 Container Threats and Isolation", 2)
add_paragraph(
    doc,
    "Jarkas et al. organize container security research into vulnerabilities, attacks, and defenses across the container "
    "lifecycle [1]. That taxonomy motivates CodeShield's attention to identity, mounts, network exposure, and image "
    "integrity. Zhou et al. propose RContainer, which extends ARM Confidential Compute Architecture primitives to strengthen "
    "container isolation [4]. He et al. show that powerful eBPF capabilities can enable cross-container attacks and increase "
    "the blast radius of a compromised component [10]. These studies demonstrate the importance of container boundaries, "
    "but their methods focus on taxonomy, hardware isolation, or attack mechanisms rather than application-specific "
    "baseline comparison. CodeShield uses their security concerns to select observable configuration fields; it does not "
    "replace those isolation mechanisms."
)

add_heading(doc, "2 2 Misconfiguration and Network Exposure", 2)
add_paragraph(
    doc,
    "Choi et al. study misconfigured container components exposed in the wild and document the practical risk created by "
    "Internet-accessible Docker and Kubernetes services [7]. Their findings support treating unexpected database publishing "
    "and other exposure changes as security-relevant drift. Rankovic et al. combine schema validation with configuration "
    "version comparison to prevent and diagnose errors in distributed cloud applications [8]. CodeShield adopts the same "
    "general principle that a configuration difference becomes more useful when it is tied to a known expected state. Its "
    "current parser is narrower because it covers selected build and container files rather than general distributed-cloud "
    "configuration."
)

add_heading(doc, "2 3 Dependency Reachability and Supply Chains", 2)
add_paragraph(
    doc,
    "ChainFuzz examines how vulnerabilities in upstream open-source components can become exploitable through downstream "
    "software supply chains [2]. Jia et al. evaluate third-party library vulnerabilities through reachability analysis, "
    "which can reduce false alarms by distinguishing included code from code that can actually be reached [3]. These works "
    "show why a version difference should not automatically be described as an exploitable vulnerability. CodeShield "
    "records a dependency mismatch and can raise its priority when known exploitation evidence is supplied, but the current "
    "prototype does not construct call graphs or independently establish reachability."
)
add_paragraph(
    doc,
    "Stalnaker et al. study how stakeholders create, consume, maintain, and verify software bills of materials [9]. Their "
    "results identify practical limitations in SBOM quality and use. CodeShield currently compares a selected dependency "
    "version and image identity with approved values, which is less comprehensive than SBOM analysis. The literature "
    "therefore supports SBOM ingestion as future work while also indicating that the resulting data will need validation "
    "and maintenance rather than simple collection."
)

add_heading(doc, "2 4 Runtime Protection and Temporary Mitigation", 2)
add_paragraph(
    doc,
    "VulShield generates temporary trigger policies that protect vulnerable code while patches are being prepared [5]. "
    "Phoenix filters system-call sequences to protect unpatched container workloads with lower overhead than coarse "
    "blocking approaches [6]. Both systems rely on precise runtime evidence and enforce protective behavior. CodeShield "
    "currently ends at detection, prioritization, persistence, and review. Manually supplied runtime observations can be "
    "evaluated, but no autonomous collector or enforcement component is implemented."
)

add_heading(doc, "2 5 Research Gap", 2)
add_paragraph(
    doc,
    "The reviewed studies provide strong specialized methods for isolation, reachability, SBOM management, exposure "
    "measurement, and runtime protection. The narrower question addressed here is whether approved application intent and "
    "observed configuration can be represented in one model and compared through explainable, evidence-linked rules. "
    "CodeShield's contribution is the implemented correlation workflow. Its value should therefore be evaluated through "
    "correct rule behavior, usable evidence extraction, traceability, and controlled experiments rather than claims that it "
    "supersedes the specialized techniques in the literature."
)

add_heading(doc, "3 Methodology", 1)
add_heading(doc, "3 1 Research and Engineering Method", 2)
add_paragraph(
    doc,
    "The work follows a design, implementation, and controlled-evaluation method. First, the supported baseline fields and "
    "drift categories were defined. A Spring Boot prototype was then implemented to acquire evidence, normalize it into a "
    "common request, apply deterministic comparisons, calculate contextual scores, and persist the results. Unit tests, web "
    "and API smoke tests, repository inspection tests, a synthetic benchmark, and two public repository scans provide the "
    "current evidence."
)

add_heading(doc, "3 2 Implemented System Architecture", 2)
figure_p = doc.add_paragraph()
figure_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure_p.paragraph_format.keep_with_next = True
figure_p.add_run().add_picture(str(FIGURE), width=Inches(6.75))
caption = doc.add_paragraph()
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_after = Pt(9)
set_font(caption.add_run("Figure 1 Implemented CodeShield processing architecture"), size=9.5, italic=True, color=MUTED)
add_paragraph(
    doc,
    "Figure 1 shows the implemented data path. User input first passes through source validation and, for repositories, "
    "file discovery and extraction. The resulting observations are converted to the same scan request used by manual and "
    "REST submissions. Drift rules compare this request with a stored baseline, and the scoring service adds contextual "
    "weights. The orchestrator persists the scan and findings for access through the web interface and JSON endpoints."
)

add_heading(doc, "3 3 Baseline and Evidence Model", 2)
add_paragraph(
    doc,
    "An ApplicationAsset stores the approved baseline, including ownership, description, criticality, expected Java "
    "version, expected database exposure, expected container user, approved image digest, and expected dependency version. "
    "A ScanRun identifies the selected asset, execution time, source reference, status, number of findings, and maximum "
    "score. Each Finding stores its rule identifier, explanation, expected state, observed state, evidence, remediation, "
    "severity, numerical score, and review status. This separation retains the relationship between intent, an observation, "
    "and the decision produced by the rule engine."
)

add_heading(doc, "3 4 Repository Acquisition and Extraction", 2)
add_paragraph(
    doc,
    "A repository scan accepts a readable local directory or a public HTTPS clone URL supplied by the user. Remote input is "
    "validated before cloning. URLs containing credentials or fragments are rejected, and address checks reject loopback, "
    "link-local, site-local, multicast, and private IPv6 destinations. The service performs a shallow clone with configurable "
    "depth and timeout into temporary storage and removes that copy after inspection. This process reads source files without "
    "modifying the target repository."
)
add_paragraph(
    doc,
    "The inspector searches for pom.xml, build.gradle, build.gradle.kts, Dockerfile variants, compose.yaml, compose.yml, "
    "docker-compose.yaml, and docker-compose.yml. Generated directories are excluded and files above the configured size "
    "limit are skipped. Maven and Gradle patterns provide Java and Spring Boot version observations. Dockerfiles provide the "
    "Java base image version and final USER; when a Dockerfile exists without USER, the observed user is treated as root. "
    "Compose files provide database port mappings, Docker socket mounts, and image references. Loopback-only database "
    "mappings are treated as internal, while other matching mappings are treated as public. Unsupported syntax remains a "
    "known parser limitation."
)

add_heading(doc, "3 5 Processing Workflow", 2)
add_numbered(doc, [
    "Register an application baseline with ownership, criticality, and approved values.",
    "Acquire evidence from a public HTTPS repository, local path, manual form, or REST request.",
    "Validate the source and apply repository safety controls where required.",
    "Discover supported files and extract the available observations.",
    "Normalize repository and manually supplied evidence into one ScanRequest representation.",
    "Apply the six deterministic rules and create one finding for each mismatch.",
    "Calculate the contextual score, map it to a severity, and persist the scan and findings.",
    "Present evidence, remediation, status, and scan history through the web interface and REST API.",
])

add_table(
    doc,
    "Table 1 Implemented drift rules",
    ["Rule", "Condition", "Default base risk"],
    [
        ["CS RUNTIME 001", "Observed Java version differs from the approved version", "30"],
        ["CS NETWORK 001", "An internally approved database is observed as public", "52"],
        ["CS CONTAINER 001", "Observed container user differs from the approved user", "42 for root  25 otherwise"],
        ["CS CONTAINER 002", "Docker socket is mounted inside the application container", "58"],
        ["CS SUPPLY 001", "Observed image digest differs from the approved digest or is unpinned", "38"],
        ["CS DEPENDENCY 001", "Observed dependency version differs from the approved version", "48 with known exploitation  24 otherwise"],
    ],
    [1.6, 4.2, 1.05],
    [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    font_size=9.1,
)

add_heading(doc, "3 6 Contextual Risk Scoring", 2)
add_paragraph(
    doc,
    "RiskScoringService calculates a transparent priority score rather than an exploit probability. All values can be "
    "overridden in application configuration or environment variables. The default calculation is shown below.",
)
formula = doc.add_paragraph()
formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
formula.paragraph_format.space_before = Pt(5)
formula.paragraph_format.space_after = Pt(7)
set_font(formula.add_run("risk score = min 100  base rule risk + criticality weight + exposure weight + exploitation weight"), name="Courier New", size=9.5, bold=True)
add_paragraph(
    doc,
    "The default criticality weights are 0 for Low, 8 for Medium, 15 for High, and 22 for Critical. Internet exposure adds "
    "15 and known exploitation evidence adds 18. Scores below 40 are Low, scores from 40 to 64 are Medium, scores from 65 "
    "to 84 are High, and scores from 85 to 100 are Critical. The same observed mismatch can therefore receive a different "
    "priority when its operational context changes."
)

add_heading(doc, "3 7 Persistence Interfaces and Reproducibility", 2)
add_paragraph(
    doc,
    "ScanOrchestrator coordinates rule execution, score calculation, and persistence. Spring MVC controllers provide the "
    "dashboard, asset and scan forms, history, and finding workflow. REST controllers expose the corresponding JSON "
    "operations. H2 provides local file persistence in the current implementation. A fresh installation begins without "
    "seeded assets or findings, which prevents demonstration records from being confused with user data."
)
add_paragraph(
    doc,
    "The rules are deterministic: the same baseline, observation, configuration, and context produce the same results. The "
    "benchmark workbook records its generation version and checksum so the controlled cases can be regenerated and audited. "
    "Repository clone depth, timeout, and file-size limits are configuration settings. These controls support repeatability, "
    "although independent execution of the full benchmark is still required before accuracy metrics can be reported."
)

add_heading(doc, "4 Experimental Results", 1)
add_heading(doc, "4 1 Automated Tests", 2)
add_paragraph(
    doc,
    "The Maven test suite completed 11 tests with no failures, errors, or skipped tests. The tests verify representative "
    "behavior in the web layer, rule engine, scoring service, and repository inspector. Table 2 summarizes the executed "
    "tests. Passing tests establish that the tested implementation paths behave as specified; they do not by themselves "
    "establish performance or generalization to arbitrary repositories."
)
add_table(
    doc,
    "Table 2 Verified automated test results",
    ["Test area", "Tests", "Verified result"],
    [
        ["Web and API smoke tests", "4", "Dashboard, asset API, repository scan page, and history page tests passed"],
        ["Drift rule tests", "2", "All six categories were detected in a constructed drift case and none in a matching case"],
        ["Risk scoring tests", "2", "Context modifiers, score cap, and severity boundaries passed"],
        ["Repository inspection tests", "3", "Evidence extraction and invalid source rejection passed"],
        ["Total", "11", "11 passed  0 failures  0 errors  0 skipped"],
    ],
    [2.0, 0.65, 4.2],
    [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
)

add_heading(doc, "4 2 Benchmark Dataset", 2)
add_paragraph(
    doc,
    "The benchmark contains 120 deterministic cases derived from the supported rule specification. Twenty-four cases are "
    "compliant and 96 contain one or more expected drift findings. Together, the cases contain 181 expected findings and "
    "are divided into 84 training, 18 validation, and 18 test cases. The split is useful for managing evaluation artifacts, "
    "but the project does not train a model."
)
add_table(
    doc,
    "Table 3 Benchmark composition",
    ["Measure", "Value"],
    [
        ["Dataset version", "CSB 2026 001"],
        ["Total cases", "120"],
        ["Compliant cases", "24"],
        ["Drift cases", "96"],
        ["Expected findings", "181"],
        ["Split", "84 train  18 validation  18 test"],
        ["Highest severity", "None 24  Low 6  Medium 31  High 33  Critical 26"],
        ["Rule frequency", "Runtime 31  Network 31  User 29  Socket 30  Image 31  Dependency 29"],
    ],
    [2.2, 4.65],
    [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
)
add_paragraph(
    doc,
    "These counts describe specification-derived ground truth, not measured predictions. They demonstrate that the dataset "
    "covers all six rule categories and a range of severity outcomes. Precision, recall, F1 score, false-positive rate, and "
    "statistical significance remain unreported because the cases have not yet been executed independently against the "
    "application and compared with their expected outputs."
)

add_heading(doc, "4 3 Public Repository Scans", 2)
add_paragraph(
    doc,
    "Two public repositories were scanned to exercise the end-to-end clone, extraction, normalization, rule, score, and "
    "persistence path. Spring Petclinic completed with five findings and a maximum score of 82. The recorded differences "
    "covered Java version, database exposure, container user, image digest, and dependency version. The "
    "ShauryaRawat27 Adaptive LLM scan completed with two findings and a maximum score of 72, covering container user and an "
    "unpinned image reference."
)
add_table(
    doc,
    "Table 4 Recorded repository scan observations",
    ["Repository", "Findings", "Maximum score", "Interpretation"],
    [
        ["spring projects spring petclinic", "5", "82", "Five differences from the selected stored baseline"],
        ["ShauryaRawat27 Adaptive LLM", "2", "72", "Two differences from the selected stored baseline"],
    ],
    [2.15, 0.75, 1.05, 2.9],
    [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
)
add_paragraph(
    doc,
    "The repository observations are preliminary integration results. The findings depend on the chosen baseline, the "
    "supported file patterns, and the configuration present at scan time. They must not be interpreted as proof that the "
    "repositories contain five or two exploitable vulnerabilities. Unsupported configuration is outside the current scan "
    "scope."
)

add_heading(doc, "4 4 Clean Installation Smoke Test", 2)
add_paragraph(
    doc,
    "A packaged JAR was started against an empty in-memory database. The dashboard, manual scan page, and repository scan "
    "page returned HTTP 200, and GET /api/assets returned an empty array. This result confirms that application startup and "
    "the tested pages do not depend on preloaded sample assets or hardcoded findings. It also separates a clean installation "
    "from the persistent local database used for earlier development scans."
)

add_heading(doc, "4 5 Remaining Evaluation", 2)
add_paragraph(
    doc,
    "Stronger empirical claims require additional experiments. All 120 cases should be executed through the application and "
    "compared with the expected rule identifiers to calculate per-rule true positives, false positives, false negatives, "
    "precision, recall, and F1 score. Repository scan time and extracted file counts should be measured across a larger, "
    "documented corpus. Paired cases that change one context variable at a time can quantify the scoring model, and "
    "independent reviewers can assess whether evidence and remediation remain understandable. Repeating scans after "
    "controlled remediation would test whether findings close consistently."
)
add_table(
    doc,
    "Table 5 Accuracy measures reserved for future execution",
    ["Rule", "TP", "FP", "FN", "Precision", "Recall", "F1"],
    [
        ["Runtime", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
        ["Network", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
        ["Container user", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
        ["Docker socket", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
        ["Image digest", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
        ["Dependency", "TBD", "TBD", "TBD", "TBD", "TBD", "TBD"],
    ],
    [1.55, 0.55, 0.55, 0.55, 1.2, 1.2, 1.2],
    [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 6,
    font_size=8.8,
)

add_heading(doc, "5 Conclusions", 1)
add_paragraph(
    doc,
    "CodeShield demonstrates a functional path from a registered application baseline and user-supplied evidence to "
    "explainable, scored, and persisted security drift findings. The implementation accepts public HTTPS repositories, "
    "local directories, browser input, and REST requests. It normalizes supported observations, executes six deterministic "
    "rules, applies configurable contextual weights, and retains the evidence required to review each result. Eleven passing "
    "tests, the controlled benchmark definition, two completed public repository scans, and the clean-start smoke test "
    "support the conclusion that the implemented workflow is operational for its defined scope."
)
add_paragraph(
    doc,
    "The available evidence does not support broader claims. Repository configuration is not equivalent to live runtime "
    "state. The parser covers selected Java and container files, and the benchmark labels originate from the same rule "
    "specification as the application. No independent precision or recall measurement has been completed. The prototype also "
    "lacks authentication, private repository credentials, autonomous collectors, vulnerability-intelligence integration, "
    "SBOM analysis, Kubernetes support, automatic remediation, and production-grade persistence. Its scores are priorities, "
    "not calibrated exploit probabilities."
)
add_paragraph(
    doc,
    "The next evaluation step is to execute all 120 benchmark cases and report per-rule accuracy measures. A broader public "
    "repository corpus can then be used to measure extraction coverage and scan time. Engineering work should add read-only "
    "authenticated runtime collectors, SBOM and vulnerability sources, access control, and PostgreSQL persistence. These "
    "changes would extend the evidence available to the same baseline comparison model while preserving the system's "
    "deterministic and explainable behavior."
)

add_heading(doc, "References", 1)
references = [
    ("[1] O. Jarkas, R. Ko, N. Dong, and R. Mahmud, A Container Security Survey Exploits Attacks and Defenses, ACM Computing Surveys, vol. 57, no. 7, Art. 170, 2025.", "https://doi.org/10.1145/3715001"),
    ("[2] P. Deng, L. Zhang, Y. Meng, et al., ChainFuzz Exploiting Upstream Vulnerabilities in Open Source Supply Chains, in 34th USENIX Security Symposium, 2025.", "https://www.usenix.org/conference/usenixsecurity25/presentation/deng"),
    ("[3] Z. Jia, C. Yang, P. Feng, et al., Impact Assessment of Third Party Library Vulnerabilities through Vulnerability Reachability Analysis, Computers and Security, vol. 157, Art. 104546, 2025.", "https://doi.org/10.1016/j.cose.2025.104546"),
    ("[4] Q. Zhou, W. Cao, X. Jia, et al., RContainer A Secure Container Architecture through Extending ARM CCA Hardware Primitives, in NDSS Symposium, 2025.", "https://doi.org/10.14722/ndss.2025.240328"),
    ("[5] Y. Li, C. Zhang, J. Zhu, et al., VulShield Protecting Vulnerable Code Before Deploying Patches, in NDSS Symposium, 2025.", "https://doi.org/10.14722/ndss.2025.240298"),
    ("[6] H. Kermabon-Bobinnec, Y. Jarraya, L. Wang, et al., Phoenix Surviving Unpatched Vulnerabilities via Accurate and Efficient Filtering of Syscall Sequences, in NDSS Symposium, 2024.", "https://doi.org/10.14722/ndss.2024.24582"),
    ("[7] D. Choi, H. Seo, K. Kim, et al., Uncovering Threats in Container Systems A Study on Misconfigured Container Components in the Wild, IEEE Access, vol. 12, pp. 192931 to 192945, 2024.", "https://doi.org/10.1109/ACCESS.2024.3514751"),
    ("[8] T. Rankovic, F. Siljic, J. Tomic, G. Sladic, and M. Simic, Misconfiguration Prevention and Error Cause Detection for Distributed Cloud Applications, in IEEE SISY, 2024.", "https://doi.org/10.1109/SISY62279.2024.10737513"),
    ("[9] T. Stalnaker, N. Wintersgill, O. Chaparro, et al., BOMs Away Inside the Minds of Stakeholders, in ICSE, 2024.", "https://doi.org/10.1145/3597503.3623347"),
    ("[10] Y. He, R. Guo, Y. Xing, et al., Cross Container Attacks The Bewildered eBPF on Clouds, in 32nd USENIX Security Symposium, 2023.", "https://www.usenix.org/conference/usenixsecurity23/presentation/he"),
]
for citation, url in references:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    set_font(p.add_run(citation + "  "), size=10)
    add_hyperlink(p, "Verified source", url)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
