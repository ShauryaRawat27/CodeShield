from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/shauryarawat/Documents/ChatGPT/fog")
OUT_DIR = ROOT / "outputs" / "codeshield-deliverables"
ARCHITECTURE = OUT_DIR / "CodeShield_System_Architecture.png"
OUTPUT = OUT_DIR / "CodeShield_Short_Project_Report.docx"

BLACK = "000000"
TEAL = "0F5965"
TEAL_DARK = "123F4A"
PALE_TEAL = "E8F3F4"
PALE_BLUE = "F3F7F8"
PALE_GOLD = "FFF8E7"
GRAY = "5E6C70"
LIGHT_GRAY = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_font(run, name="Liberation Sans", size=None, color=BLACK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.12):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_body(doc, text, *, bold_lead=None, after=6, keep_with_next=False):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        run = p.add_run(text)
        set_font(run)
    set_paragraph_spacing(p, after=after)
    p.paragraph_format.keep_with_next = keep_with_next
    return p


def add_bullet(doc, text, level=0, after=3):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    set_font(run, size=10.5)
    set_paragraph_spacing(p, after=after, line=1.08)
    return p


def add_numbered(doc, text, after=3):
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    set_font(run, size=10.5)
    set_paragraph_spacing(p, after=after, line=1.08)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, size=15 if level == 1 else 12, bold=True, color=BLACK)
    return p


def add_table(doc, headers, rows, widths, *, header_fill=TEAL_DARK):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])

    for idx, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[idx]
        set_cell_width(cell, width)
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(header)
        set_font(run, size=9.5, color=WHITE, bold=True)
        set_paragraph_spacing(p, after=0, line=1.0)

    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        for cidx, (value, width) in enumerate(zip(row_data, widths)):
            cell = row.cells[cidx]
            set_cell_width(cell, width)
            set_cell_shading(cell, WHITE if ridx % 2 == 0 else PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            run = p.add_run(value)
            set_font(run, size=9.3)
            set_paragraph_spacing(p, after=0, line=1.05)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    return table


def add_page_number(paragraph):
    run = paragraph.add_run("Page ")
    set_font(run, size=8.5, color=GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = p_bdr.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            p_bdr.append(node)
        node.set(qn("w:val"), "nil")


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Liberation Sans"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(9)

    for name, size, before, after in (
        ("Heading 1", 15, 13, 6),
        ("Heading 2", 12, 9, 4),
    ):
        style = styles[name]
        style.font.name = "Liberation Sans"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = "Liberation Sans"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(BLACK)


def add_image_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def build_report():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if not ARCHITECTURE.exists():
        raise FileNotFoundError(f"Architecture image not found: {ARCHITECTURE}")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.28)
    configure_styles(doc)

    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(
        Inches(7.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES
    )
    left_run = footer.add_run("CodeShield Project Report")
    set_font(left_run, size=8.5, color=GRAY)
    tab_run = footer.add_run("\t")
    set_font(tab_run, size=8.5, color=GRAY)
    add_page_number(footer)
    set_paragraph_spacing(footer, after=0, line=1.0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    remove_paragraph_borders(title)
    title_run = title.add_run(
        "CodeShield Context Aware SecOps Platform for Detecting Application Infrastructure Security Drift"
    )
    set_font(title_run, name="Liberation Sans", size=22, bold=True, color=BLACK)

    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("Short Project Report")
    set_font(subtitle_run, name="Liberation Sans", size=13, bold=True, color=TEAL)
    set_paragraph_spacing(subtitle, after=2)

    meta = doc.add_paragraph()
    meta_run = meta.add_run("Cybersecurity and DevSecOps  |  October 2026")
    set_font(meta_run, size=9.5, color=GRAY)
    set_paragraph_spacing(meta, after=18)

    add_heading(doc, "Executive Summary", 1)
    add_body(
        doc,
        "CodeShield is a security platform that detects differences between an application's approved baseline and the configuration found in its repository or deployed environment. It combines repository evidence with supplied observations from containers, dependencies, hosts, and networks. The platform produces prioritized findings that state what was expected, what was observed, why the difference matters, and how to correct it.",
    )
    add_body(
        doc,
        "The implemented application combines a read-only local repository scanner, manual and API-based runtime evidence input, deterministic drift rules, and transparent risk scoring. Findings retain their evidence source and can be marked open, resolved, or accepted. The system runs as a Spring Boot web application with an embedded database and Docker support.",
        after=9,
    )

    overview_rows = [
        ("Primary goal", "Detect security drift between application intent and runtime reality"),
        ("Main inputs", "Maven, Gradle, Dockerfile and Compose files plus supplied runtime observations"),
        ("Main output", "Explainable findings with evidence, risk score and remediation guidance"),
        ("Implemented platform", "Java 21, Spring Boot 3.3, Thymeleaf, Spring Data JPA and H2"),
    ]
    add_table(doc, ["Area", "Project definition"], overview_rows, [1.55, 5.35])

    add_heading(doc, "Problem and Objectives", 1)
    add_body(
        doc,
        "An application can be secure in source control but unsafe after deployment. A database intended for an internal network may become publicly reachable, a container may run as root, or the deployed library version may differ from the approved SBOM. These changes may result from configuration edits, emergency fixes, image replacement, or infrastructure maintenance.",
    )
    add_body(doc, "CodeShield has four practical objectives:", keep_with_next=True)
    add_bullet(doc, "Store an approved application baseline covering versions, exposure, image identity and container user.")
    add_bullet(doc, "Extract evidence from Maven, Gradle, Dockerfile and Docker Compose configuration without modifying the repository.")
    add_bullet(doc, "Apply explainable rules to identify meaningful mismatches between the two models.")
    add_bullet(doc, "Prioritize each mismatch using exposure, reachability, exploitability and asset importance.", after=7)

    add_heading(doc, "How CodeShield Differs", 1)
    comparison_rows = [
        ("Analysis scope", "Usually one layer or data source", "Connects application, container, host, dependency and network evidence"),
        ("Decision basis", "Reports a vulnerability or misconfiguration", "Checks whether the observed state violates application intent"),
        ("Prioritization", "Often dominated by severity scores", "Adds exposure, reachability, privilege, asset importance and compensating controls"),
        ("Explanation", "Describes the isolated alert", "Records expected state, observed state, evidence and recommended correction"),
        ("Ongoing use", "Runs periodic independent scans", "Tracks drift and verifies remediation through re-scanning"),
    ]
    add_table(
        doc,
        ["Dimension", "Conventional scanner", "CodeShield"],
        comparison_rows,
        [1.25, 2.55, 3.10],
    )

    add_heading(doc, "Illustrative Finding", 2)
    add_body(
        doc,
        "A configuration file declares Redis as an internal service, while the Docker and network collectors observe port 6379 published on all host interfaces. CodeShield correlates these facts and reports a trust-boundary violation. The finding identifies the expected private access, the observed public exposure, the affected asset, the supporting evidence, and the recommended action to remove the host mapping or restrict network access.",
    )

    add_heading(doc, "Project Scope", 2)
    add_body(
        doc,
        "CodeShield is designed for authorized security monitoring and complements existing source, dependency and container scanners. It can consume their results as evidence, while its own responsibility is to correlate that evidence with the deployed environment.",
    )
    add_bullet(doc, "Inspect only registered applications, hosts and container environments.")
    add_bullet(doc, "Use read-only collection wherever possible and preserve evidence timestamps.")
    add_bullet(doc, "Detect and report drift without automatically changing production systems.")
    add_bullet(doc, "Require a human decision for exceptions and disruptive remediation.", after=6)

    doc.add_page_break()
    add_heading(doc, "System Architecture", 1)
    add_body(
        doc,
        "The architecture separates evidence collection from analysis and decision-making. Read-only collectors and authenticated agents gather facts. The Spring Boot control plane schedules work, manages access, stores evidence, and exposes findings. The correlation engine compares normalized intent and runtime models before the risk engine prioritizes the result.",
        after=7,
    )

    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    image_run = image_paragraph.add_run()
    shape = image_run.add_picture(str(ARCHITECTURE), width=Inches(7.02))
    add_image_alt_text(
        shape,
        "CodeShield system architecture",
        "Evidence sources feed a Spring Boot control plane and analyzer workers. The resulting application intent and runtime reality models enter a correlation engine, followed by risk scoring, storage, and findings.",
    )

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = caption.add_run("Figure 1  CodeShield system architecture")
    set_font(cap_run, size=9, color=GRAY, italic=True)
    set_paragraph_spacing(caption, before=3, after=7, line=1.0)

    component_rows = [
        ("Evidence collection", "Reads source and deployment artifacts; queries Docker; receives host, OS and network observations."),
        ("Normalized models", "Converts raw inputs into comparable application intent and runtime reality facts."),
        ("Correlation engine", "Applies deterministic drift rules and preserves the evidence behind each decision."),
        ("Risk engine", "Combines severity with exposure, criticality, mismatch, reachability, privilege and controls."),
        ("Findings workflow", "Supports dashboards, reports, ownership, exceptions, remediation status and re-scans."),
    ]
    add_table(doc, ["Layer", "Responsibility"], component_rows, [1.65, 5.25])

    add_heading(doc, "Context Aware Risk Logic", 1)
    add_body(
        doc,
        "The first version should use transparent rules rather than machine learning. A finding receives more weight when the affected asset is important, the service is externally exposed, the vulnerable code is reachable, or the runtime has excessive privileges. Verified controls such as network restrictions or compensating authentication may reduce the final score. CVSS remains an input, but it does not determine priority by itself.",
    )

    rule_rows = [
        ("Java 21 required", "Java 17 running", "Runtime compatibility and patching risk"),
        ("Database intended to be internal", "Database port publicly reachable", "Network trust-boundary violation"),
        ("Non-root container expected", "Container runs as root", "Excessive runtime privilege"),
        ("Approved image digest", "Different digest running", "Deployment integrity drift"),
        ("Dependency version X declared", "Version Y deployed", "Dependency or SBOM drift"),
    ]
    add_table(doc, ["Expected state", "Observed state", "Result"], rule_rows, [2.1, 2.15, 2.65])

    add_heading(doc, "Implemented Scope", 1)
    add_body(
        doc,
        "The current application implements the correlation workflow with focused evidence sources and explainable rules. It avoids automated remediation and machine learning so every finding can be traced to an explicit comparison.",
    )
    mvp_items = [
        "Register application baselines with ownership, criticality and approved configuration.",
        "Inspect local repositories for Maven, Gradle, Dockerfile and Docker Compose evidence.",
        "Accept runtime observations through a browser form or JSON API.",
        "Detect Java, dependency, exposure, container-user, Docker-socket and image-integrity drift.",
        "Score findings using rule risk, application criticality, Internet exposure and exploit context.",
        "Display evidence, expected and observed states, remediation, workflow status and scan history.",
        "Persist data in H2 and run locally or through Docker Compose.",
    ]
    for item in mvp_items:
        add_numbered(doc, item)

    add_heading(doc, "Security Design", 1)
    security_rows = [
        ("Repository access", "Read only supported configuration files and never modify the target repository."),
        ("Path handling", "Normalize and validate repository paths before inspection."),
        ("File scope", "Ignore generated directories and reject configuration files larger than one megabyte."),
        ("Runtime data", "Record the evidence source and keep repository findings separate from live-host claims."),
        ("Remediation", "Provide guidance and workflow status without changing the scanned system."),
        ("Deployment", "Run the application container as a non-root user and mount scanned repositories read only."),
    ]
    add_table(doc, ["Boundary", "Implemented control"], security_rows, [1.45, 5.45])

    add_heading(doc, "Evaluation Plan", 1)
    add_body(
        doc,
        "Evaluation should use controlled deployments with seeded drift cases and a known ground truth. Each case should represent a realistic mismatch, such as a publicly exposed internal service, an unapproved image digest, an incorrect runtime version, or an unnecessary privilege.",
    )
    add_bullet(doc, "Measure precision, recall and false-positive rate for each drift category.")
    add_bullet(doc, "Compare alert quality before and after application context is added.")
    add_bullet(doc, "Measure dependency matching and network exposure classification accuracy.")
    add_bullet(doc, "Record scan duration, agent overhead, stale-evidence behavior and re-scan consistency.", after=7)

    add_heading(doc, "Current Result", 1)
    add_body(
        doc,
        "CodeShield now scans real local repositories, converts supported configuration into comparable evidence, and applies the same drift and risk engine used for submitted runtime observations. Automated tests cover repository extraction, rule detection, risk scoring, web pages and JSON endpoints. The accompanying literature review workbook provides the research foundation for container security, software supply-chain risk, vulnerability prioritization, SBOM analysis, runtime monitoring and configuration drift.",
    )

    add_heading(doc, "Project Deliverables", 1)
    add_bullet(doc, "Runnable Spring Boot source code with web interface, REST API, database and Docker support.")
    add_bullet(doc, "CodeShield System Architecture in PNG and editable SVG formats.")
    add_bullet(doc, "CodeShield Literature Review with ten recent and verified papers in Excel format.")
    add_bullet(doc, "CodeShield README with detailed architecture, modules, security boundaries and implementation scope.")

    # Stable metadata without personal details.
    props = doc.core_properties
    props.title = "CodeShield Context Aware SecOps Platform for Detecting Application Infrastructure Security Drift"
    props.subject = "Short project report"
    props.keywords = "CodeShield, SecOps, security drift, Docker, Spring Boot, vulnerability management"
    props.author = ""
    props.last_modified_by = ""

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
